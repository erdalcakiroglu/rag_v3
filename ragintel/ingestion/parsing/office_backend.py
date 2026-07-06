"""xlsx (openpyxl) ve txt (charset-normalizer) parse — backend'den bağımsız.

Bu iki tip Docling'e gitmez (İP-2 kuralı); her zaman bu fonksiyonlarla işlenir.
"""

from __future__ import annotations

from .parsed_document import Page, ParsedDocument, Table
from .text_utils import detect_language, flatten_table


def parse_xlsx(path: str) -> ParsedDocument:
    """Her sheet -> bir Table (sheet_name'li). İçerik tabloda; gövde boş."""
    import openpyxl

    pd = ParsedDocument()
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    try:
        for idx, ws in enumerate(wb.worksheets):
            rows: list[list] = []
            for row in ws.iter_rows(values_only=True):
                if row is None:
                    continue
                # Tamamen boş satırları atla.
                if all(c is None for c in row):
                    continue
                rows.append(list(row))
            if not rows:
                continue
            pd.tables.append(Table(
                index=idx,
                data=rows,
                flattened_text=flatten_table(rows),
                sheet_name=ws.title,
            ))
    finally:
        wb.close()

    if not pd.tables:
        pd.warn("xlsx: veri içeren sheet bulunamadı")
    pd.language = detect_language(
        "\n".join(t.flattened_text for t in pd.tables)
    )
    return pd


def parse_txt(path: str) -> ParsedDocument:
    """charset-normalizer ile decode; paragraflara böl (tek mantıksal sayfa)."""
    from charset_normalizer import from_path

    pd = ParsedDocument()
    best = from_path(path).best()
    if best is None:
        pd.warn("txt: encoding tespit edilemedi")
        pd.pages.append(Page(page_no=1, text_blocks=[]))
        return pd

    text = str(best)
    if best.encoding:
        pd.warn(f"txt: encoding={best.encoding}")

    blocks = [b.strip() for b in text.split("\n\n") if b.strip()]
    if not blocks and text.strip():
        blocks = [text.strip()]
    pd.pages.append(Page(page_no=1, text_blocks=blocks))
    pd.language = detect_language(text)
    return pd
