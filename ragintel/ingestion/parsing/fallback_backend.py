"""Yedek parse backend'i: pdf=pymupdf, docx=python-docx.

Docling kurulu olmadığında veya hızlı/offline (test) çalıştırmada kullanılır.
Tablolar gövde metninden AYRILIR (İP-2 kuralı): pdf'te tablo bbox'ları ile
kesişen metin blokları gövdeye alınmaz. Sayfa numaraları korunur; section
başlıkları pdf'te font-boyutu, docx'te stil (Heading) ile tespit edilir.

OCR: bu backend OCR yapmaz (gerçek OCR Docling'e aittir). ocr=True gelirse
uyarı eklenir; adaptörün OCR-fallback ORKESTRASYONU backend'den bağımsızdır.
"""

from __future__ import annotations

from collections import Counter

from .parsed_document import Figure, Page, ParsedDocument, Section, Table
from .text_utils import detect_language, flatten_table


class FallbackBackend:
    name = "fallback"

    def supports(self, file_type: str) -> bool:
        return file_type in ("pdf", "docx")

    def parse(self, path: str, file_type: str, *, ocr: bool = False) -> ParsedDocument:
        if file_type == "pdf":
            return _parse_pdf(path, ocr=ocr)
        if file_type == "docx":
            return _parse_docx(path)
        raise ValueError(f"FallbackBackend desteklemiyor: {file_type}")


# --------------------------------------------------------------------------
# PDF (pymupdf / fitz)
# --------------------------------------------------------------------------
def _parse_pdf(path: str, *, ocr: bool = False) -> ParsedDocument:
    import fitz

    pd = ParsedDocument()
    doc = fitz.open(path)
    try:
        dominant = _dominant_font_size(doc)
        heading_min = dominant * 1.15 if dominant else 0
        offset = 0
        tbl_index = 0
        fig_index = 0

        for pno in range(len(doc)):
            page = doc[pno]
            page_no = pno + 1

            # 1) Tablolar (ayrı; gövdeye karışmaz).
            table_rects = []
            try:
                found = page.find_tables()
                tabs = getattr(found, "tables", found)
            except Exception as exc:  # tablo motoru sayfada patlarsa gövde devam
                tabs = []
                pd.warn(f"s{page_no}: tablo tespiti hatası ({exc})")
            for t in tabs:
                rows = t.extract()
                pd.tables.append(Table(
                    index=tbl_index, data=rows,
                    flattened_text=flatten_table(rows), page_no=page_no,
                ))
                tbl_index += 1
                table_rects.append(fitz.Rect(t.bbox))

            # 2) Metin blokları + section + figürler.
            text_blocks: list[str] = []
            data = page.get_text("dict")
            for block in data.get("blocks", []):
                if block.get("type") == 1:  # görsel
                    pd.figures.append(Figure(index=fig_index, page_no=page_no))
                    fig_index += 1
                    continue
                brect = fitz.Rect(block["bbox"])
                if any(brect.intersects(tr) for tr in table_rects):
                    continue  # tablo alanı -> gövdeye alma
                for line in block.get("lines", []):
                    spans = line.get("spans", [])
                    text = "".join(s["text"] for s in spans).strip()
                    if not text:
                        continue
                    size = max((s["size"] for s in spans), default=0)
                    if heading_min and size >= heading_min and len(text) < 80:
                        pd.sections.append(Section(
                            title=text,
                            level=_heading_level(size, dominant),
                            page_start=page_no,
                            char_span=(offset, offset + len(text)),
                        ))
                    text_blocks.append(text)
                    offset += len(text) + 1

            pd.pages.append(Page(page_no=page_no, text_blocks=text_blocks))

        if ocr:
            pd.warn("fallback backend OCR desteklemiyor (gerçek OCR: Docling)")
        pd.language = detect_language(pd.body_text)
    finally:
        doc.close()
    return pd


def _dominant_font_size(doc) -> float:
    """Belgedeki en yaygın span font boyutu (gövde metni referansı)."""
    sizes: Counter = Counter()
    for pno in range(min(len(doc), 10)):  # ilk 10 sayfa yeterli örnek
        for block in doc[pno].get_text("dict").get("blocks", []):
            for line in block.get("lines", []):
                for s in line.get("spans", []):
                    sizes[round(s["size"], 1)] += len(s.get("text", ""))
    return sizes.most_common(1)[0][0] if sizes else 0.0


def _heading_level(size: float, dominant: float) -> int:
    if not dominant:
        return 1
    ratio = size / dominant
    if ratio >= 1.6:
        return 1
    if ratio >= 1.3:
        return 2
    return 3


# --------------------------------------------------------------------------
# DOCX (python-docx)
# --------------------------------------------------------------------------
def _parse_docx(path: str) -> ParsedDocument:
    from docx import Document

    pd = ParsedDocument()
    doc = Document(path)
    page = Page(page_no=1)
    offset = 0

    for para in doc.paragraphs:
        text = para.text.strip()
        if not text:
            continue
        style = (para.style.name if para.style else "") or ""
        if style.startswith("Heading") or style == "Title":
            pd.sections.append(Section(
                title=text,
                level=_docx_heading_level(style),
                page_start=1,
                char_span=(offset, offset + len(text)),
            ))
        page.text_blocks.append(text)
        offset += len(text) + 1

    for ti, tbl in enumerate(doc.tables):
        rows = [[cell.text for cell in row.cells] for row in tbl.rows]
        pd.tables.append(Table(
            index=ti, data=rows,
            flattened_text=flatten_table(rows), page_no=1,
        ))

    pd.pages.append(page)
    pd.language = detect_language(pd.body_text)
    return pd


def _docx_heading_level(style: str) -> int:
    if style == "Title":
        return 1
    # "Heading 2" -> 2
    parts = style.split()
    if len(parts) == 2 and parts[1].isdigit():
        return int(parts[1])
    return 1
