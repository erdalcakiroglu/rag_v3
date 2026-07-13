"""İP-10 Orchestrator — uçtan uca akış + durum makinesi (canlı DB, mock embed).

Kabul: temsili korpus tek komutla uçtan uca COMPLETED; kill -9 sonrası kaldığı
yerden devam + çift işleme yok. Ek: RETRY, REPROCESS (parse/intake başlangıç),
crash-recovery, EmbeddingBackendError -> PENDING (retry_count artmaz).
"""

from __future__ import annotations

import pytest

from ragintel.config.loader import load_config
from ragintel.database.config_store import make_db_reader
from ragintel.ingestion import Orchestrator
from ragintel.ingestion.chunking import WordTokenCounter
from ragintel.ingestion.embedding import OllamaEmbedder
from ragintel.ingestion.parsing import get_backend
from ragintel.ingestion.scanner import FolderScanner
from tests import _corpus_parse as cp
from tests import _ollama_mock as om

pytestmark = pytest.mark.db

TEST_SCOPE = "__ip10_test__"


def _cfg(db, **ingestion_over):
    base = make_db_reader(db)

    def reader():
        groups = dict(base())
        ing = dict(groups.get("ingestion", {}))
        ing.update(ingestion_over)
        groups["ingestion"] = ing
        return groups
    return load_config(db_reader=reader)


def _orch(db, tmp_path, *, embedder=None, config=None):
    embedder = embedder or OllamaEmbedder("http://mock", client=om.make_client(om.ok_handler()))
    return Orchestrator(
        db, config=config or _cfg(db),
        embedder=embedder,
        parse_backend=get_backend("fallback"),
        token_counter=WordTokenCounter(),
        storage_root=str(tmp_path / "storage"),
    )


def _build_corpus(root):
    import os
    from docx import Document

    os.makedirs(root, exist_ok=True)
    marker = os.path.basename(root)
    cp.make_text_pdf(
        os.path.join(root, f"a_{marker}.pdf"),
        title=f"Rapor A {marker}",
        body=(f"Karbon vergisi govde A {marker}. " * 6),
    )
    cp.make_text_pdf(
        os.path.join(root, f"b_{marker}.pdf"),
        title=f"Rapor B {marker}",
        body=(f"Karbon vergisi govde B {marker}. " * 6),
    )

    doc = Document()
    doc.add_heading(f"Ana Başlık {marker}", level=1)
    doc.add_paragraph(f"Karbon vergisi giriş paragrafı metni {marker}.")
    doc.add_heading(f"Alt Başlık {marker}", level=2)
    doc.add_paragraph(f"İkinci bölüm paragraf metni burada {marker}.")
    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).text = f"TbAd{marker}"
    table.cell(0, 1).text = "TbDeger"
    table.cell(1, 0).text = "TbKarbon"
    table.cell(1, 1).text = f"TbVal{marker}"
    doc.save(os.path.join(root, f"c_{marker}.docx"))

    cp.make_xlsx(os.path.join(root, f"d_{marker}.xlsx"), title=f"deg-{marker}")
    cp.make_txt(
        os.path.join(root, f"e_{marker}.txt"),
        text=f"Karbon vergisi düz metin {marker}.\n\nİkinci paragraf {marker}.",
    )
    return 5


def _status_counts(db):
    with db.connection() as conn:
        return {s: c for s, c in conn.execute(
            "SELECT status, count(*) FROM core_files WHERE doc_scope=%s GROUP BY status;",
            (TEST_SCOPE,)).fetchall()}


def _file_ids(db):
    with db.connection() as conn:
        return [r[0] for r in conn.execute(
            "SELECT file_id FROM core_files WHERE doc_scope=%s ORDER BY file_id;",
            (TEST_SCOPE,)).fetchall()]


def _chunk_count(db, file_id):
    with db.connection() as conn:
        return conn.execute("SELECT count(*) FROM core_chunks WHERE file_id=%s;",
                            (file_id,)).fetchone()[0]


@pytest.fixture(autouse=True)
def _clean(live_db):
    with live_db.connection() as conn:
        conn.execute("DELETE FROM core_files WHERE doc_scope=%s;", (TEST_SCOPE,))
    yield
    with live_db.connection() as conn:
        conn.execute("DELETE FROM core_files WHERE doc_scope=%s;", (TEST_SCOPE,))


# --- Uçtan uca korpus -> COMPLETED -------------------------------------------
def test_end_to_end_corpus_completed(live_db, tmp_path):
    n = _build_corpus(str(tmp_path / "corpus"))
    orch = _orch(live_db, tmp_path)
    scan = orch.scan(str(tmp_path / "corpus"), doc_scope=TEST_SCOPE)
    assert scan.total == n and scan.summary()["created"] == n

    orch.run()
    # NOT (izolasyon): orch.run() scope'suz çalışır — PENDING kuyruğunu TÜM
    # doc_scope'lar için işler. Paylaşılan canlı DB'de aynı anda başka
    # scope'lardan bekleyen/işlenen kayıtlar varsa run()'ın dönüş sözlüğü
    # (COMPLETED/FAILED toplamı) bu testten BAĞIMSIZ olarak değişebilir —
    # bu yüzden dönüş değerine değil, SADECE bu testin ürettiği doc_scope'a
    # scope'lu DB durumuna güveniyoruz: n dosyanın TÜMÜ COMPLETED, başka
    # durumda (FAILED/PENDING/PROCESSING) kalan yok.
    assert _status_counts(live_db) == {"COMPLETED": n}
    # chunk + vektör üretildi
    with live_db.connection() as conn:
        ch = conn.execute(
            "SELECT count(*) FROM core_chunks c JOIN core_files f USING(file_id) "
            "WHERE f.doc_scope=%s;", (TEST_SCOPE,)).fetchone()[0]
        vec = conn.execute(
            "SELECT count(*) FROM core_vectors v JOIN core_chunks c USING(chunk_id) "
            "JOIN core_files f USING(file_id) WHERE f.doc_scope=%s;",
            (TEST_SCOPE,)).fetchone()[0]
    assert ch > 0 and vec == ch


# --- kill -9 sonrası devam + çift işleme yok ---------------------------------
def test_kill9_resume_no_double_processing(live_db, tmp_path):
    _build_corpus(str(tmp_path / "corpus"))
    orch = _orch(live_db, tmp_path, config=_cfg(live_db, stuck_processing_minutes=0))
    orch.scan(str(tmp_path / "corpus"), doc_scope=TEST_SCOPE)
    orch.run()
    assert _status_counts(live_db) == {"COMPLETED": 5}

    victim = _file_ids(live_db)[0]
    before = _chunk_count(live_db, victim)
    # kill -9 simülasyonu: dosya PROCESSING'de takılı kaldı.
    with live_db.connection() as conn:
        conn.execute("UPDATE core_files SET status='PROCESSING' WHERE file_id=%s;",
                     (victim,))

    orch.run()   # açılışta recover_stuck -> victim RETRY -> yeniden işlenir
    assert _status_counts(live_db) == {"COMPLETED": 5}
    # çift işleme yok: chunk sayısı aynı (delete_file_derived ile temizlenip yazıldı)
    assert _chunk_count(live_db, victim) == before


# --- crash-recovery: PROCESSING takılı -> RETRY ------------------------------
def test_recover_stuck_moves_to_pending(live_db, tmp_path):
    _build_corpus(str(tmp_path / "corpus"))
    orch = _orch(live_db, tmp_path, config=_cfg(live_db, stuck_processing_minutes=0))
    orch.scan(str(tmp_path / "corpus"), doc_scope=TEST_SCOPE)
    fid = _file_ids(live_db)[0]
    with live_db.connection() as conn:
        conn.execute("UPDATE core_files SET status='PROCESSING' WHERE file_id=%s;", (fid,))

    stuck = orch.recover_stuck()
    assert fid in stuck
    with live_db.connection() as conn:
        st, rc = conn.execute("SELECT status, retry_count FROM core_files WHERE file_id=%s;",
                              (fid,)).fetchone()
    assert st == "PENDING" and rc == 1     # RETRY'a çekildi (retry_count++)


# --- RETRY: FAILED (retry_count<3) yeniden işlenir ---------------------------
def test_retry_failed_file(live_db, tmp_path):
    _build_corpus(str(tmp_path / "corpus"))
    orch = _orch(live_db, tmp_path)
    orch.scan(str(tmp_path / "corpus"), doc_scope=TEST_SCOPE)
    orch.run()
    fid = _file_ids(live_db)[0]
    # Bir dosyayı elle FAILED yap (retry_count=0), türevleri sil.
    with live_db.connection() as conn:
        conn.execute("DELETE FROM core_chunks WHERE file_id=%s;", (fid,))
        conn.execute("UPDATE core_files SET status='FAILED', retry_count=0 WHERE file_id=%s;",
                     (fid,))

    res = orch.retry()
    assert res.get("COMPLETED", 0) >= 1
    with live_db.connection() as conn:
        st, rc = conn.execute("SELECT status, retry_count FROM core_files WHERE file_id=%s;",
                              (fid,)).fetchone()
    assert st == "COMPLETED" and rc == 1


# --- REPROCESS (parse'dan): COMPLETED dosya yeniden işlenir -------------------
def test_reprocess_from_parse_renews_chunks(live_db, tmp_path):
    _build_corpus(str(tmp_path / "corpus"))
    orch = _orch(live_db, tmp_path)
    orch.scan(str(tmp_path / "corpus"), doc_scope=TEST_SCOPE)
    orch.run()
    fid = _file_ids(live_db)[0]
    with live_db.connection() as conn:
        old = [r[0] for r in conn.execute(
            "SELECT chunk_id FROM core_chunks WHERE file_id=%s;", (fid,)).fetchall()]

    assert orch.reprocess(fid) == "COMPLETED"
    with live_db.connection() as conn:
        new = [r[0] for r in conn.execute(
            "SELECT chunk_id FROM core_chunks WHERE file_id=%s;", (fid,)).fetchall()]
    assert old and new and set(old).isdisjoint(new)   # chunk_id'ler yenilendi


# --- REPROCESS (intake'ten): intake-FAILED (oversize) dosya kurtarılır --------
def test_reprocess_from_intake_recovers_oversize(live_db, tmp_path):
    import os
    root = str(tmp_path / "corpus")
    os.makedirs(root)
    cp.make_txt(os.path.join(root, "big.txt"), "x" * 5000)

    # Tiny limitli ayrı scanner ile intake-FAILED (raw kopyasız) üret.
    scanner = FolderScanner(live_db, config=_cfg(live_db),
                            storage_root=str(tmp_path / "storage"))
    scanner.max_bytes = 100
    scanner.scan(root, doc_scope=TEST_SCOPE)
    fid = _file_ids(live_db)[0]
    with live_db.connection() as conn:
        st = conn.execute("SELECT status FROM core_files WHERE file_id=%s;", (fid,)).fetchone()[0]
    assert st == "FAILED"     # oversize -> intake FAILED

    # Normal limitli orchestrator ile REPROCESS -> reintake'ten başlar.
    orch = _orch(live_db, tmp_path)   # varsayılan max 100MB
    assert orch.reprocess(fid) == "COMPLETED"
    assert _chunk_count(live_db, fid) > 0


# --- EmbeddingBackendError -> PENDING (dosya FAILED değil, retry_count artmaz) -
def test_embedding_backend_error_reverts_to_pending(live_db, tmp_path):
    import os
    root = str(tmp_path / "corpus")
    os.makedirs(root)
    marker = os.path.basename(root)
    cp.make_txt(
        os.path.join(root, f"a_{marker}.txt"),
        text=f"Embedding hata testi {marker}.",
    )
    orch = _orch(live_db, tmp_path,
                 embedder=OllamaEmbedder("http://mock", client=om.make_client(om.unreachable_handler())))
    orch.scan(root, doc_scope=TEST_SCOPE)
    fid = _file_ids(live_db)[0]

    from ragintel.ingestion.embedding import EmbeddingBackendError
    with pytest.raises(EmbeddingBackendError):
        orch.run()

    with live_db.connection() as conn:
        st, rc = conn.execute("SELECT status, retry_count FROM core_files WHERE file_id=%s;",
                              (fid,)).fetchone()
    assert st == "PENDING" and rc == 0     # altyapı hatası: dosya hatası değil


# --- status raporu -----------------------------------------------------------
def test_status_counts(live_db, tmp_path):
    _build_corpus(str(tmp_path / "corpus"))
    orch = _orch(live_db, tmp_path)
    orch.scan(str(tmp_path / "corpus"), doc_scope=TEST_SCOPE)
    orch.run()
    counts = orch.status()
    # NOT (vacuous-test denetimi): orch.status() scope'suz TÜM DB'yi sayar;
    # paylaşılan canlı DB'de başka kayıtlar varsa `>= 5` bu testin KENDİ
    # işlemesinden bağımsız olarak da doğru çıkabilir (yokluk/rastlantı
    # riski). Global sayaç en az bizim scope'umuz kadar olmalı diye ZAYIF
    # bir sağlık kontrolü olarak tutuyoruz; ama asıl (pozitif, scope'lu)
    # kanıt bir alt satırda: testin ürettiği n=5 dosyanın GERÇEKTEN
    # COMPLETED olduğu doğrudan doc_scope'a filtrelenmiş sorguyla kanıtlanır.
    assert counts.get("COMPLETED", 0) >= 5
    assert _status_counts(live_db) == {"COMPLETED": 5}
